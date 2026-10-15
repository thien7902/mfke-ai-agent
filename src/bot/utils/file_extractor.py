"""Extract text from files uploaded to Telegram.

`extract_text_from_file` dispatches by extension/MIME and returns text for
text/PDF/DOCX/XLSX files. For images it returns the `IMAGE_SENTINEL` string
so the caller knows to use the image pathway (vision model pre-step or OCR).
"""
import logging
from pathlib import Path
from typing import Optional

import structlog

from src.bot.utils.config import config

logger = structlog.get_logger(__name__)

# Returned by `extract_text_from_file` for image types. The handler uses this
# to decide between vision-model pre-step and OCR fallback.
IMAGE_SENTINEL = "__IMAGE__"

# Extensions we treat as plain text / code — read as UTF-8.
TEXT_EXTENSIONS = {
    ".txt", ".md", ".markdown", ".rst",
    ".py", ".js", ".ts", ".tsx", ".jsx", ".java", ".kt", ".swift", ".go",
    ".rs", ".rb", ".php", ".c", ".h", ".cpp", ".hpp", ".cc", ".cs",
    ".sh", ".bash", ".zsh", ".fish", ".ps1",
    ".json", ".yaml", ".yml", ".toml", ".ini", ".cfg", ".conf",
    ".csv", ".tsv", ".log",
    ".sql", ".graphql", ".gql",
    ".html", ".htm", ".css", ".scss", ".xml", ".svg",
    ".env", ".gitignore", ".dockerfile",
    ".tf", ".hcl",
}

IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".tiff"}


def _truncate(text: str) -> str:
    """Cap text length at config.file_upload_max_text_bytes with a notice."""
    cap = config.file_upload_max_text_bytes
    if len(text.encode("utf-8", errors="replace")) <= cap:
        return text
    # Encode then slice to honor byte budget; decode back to str.
    encoded = text.encode("utf-8", errors="replace")[:cap]
    truncated = encoded.decode("utf-8", errors="ignore")
    return truncated + "\n\n[... file truncated at %d bytes ...]" % cap


def _extract_text_file(path: Path) -> str:
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        return f.read()


def _extract_pdf(path: Path) -> str:
    from pypdf import PdfReader
    reader = PdfReader(str(path))
    parts = []
    for page in reader.pages:
        try:
            txt = page.extract_text() or ""
        except Exception as e:
            logger.warning("PDF page extraction failed", path=str(path), error=str(e))
            txt = ""
        parts.append(txt)
    return "\n\n".join(parts)


def _extract_docx(path: Path) -> str:
    import docx
    doc = docx.Document(str(path))
    parts = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                txt = cell.text.strip()
                if txt:
                    parts.append(txt)
    return "\n".join(parts)


def _extract_xlsx(path: Path) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(str(path), read_only=True, data_only=True)
    parts = []
    try:
        for sheet in wb.worksheets:
            parts.append(f"### Sheet: {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None]
                if cells:
                    parts.append("\t".join(cells))
    finally:
        wb.close()
    return "\n".join(parts)


def extract_text_from_image(path: Path) -> str:
    """OCR an image via tesseract. Returns "" if tesseract is unavailable."""
    try:
        import pytesseract
        from PIL import Image
    except ImportError as e:
        logger.warning("OCR libs missing", error=str(e))
        return ""
    try:
        return pytesseract.image_to_string(Image.open(str(path)))
    except Exception as e:
        logger.warning("OCR failed", path=str(path), error=str(e))
        return ""


def extract_text_from_file(path: Path, mime: Optional[str] = None) -> str:
    """Extract text from a downloaded file.

    Returns:
        - Extracted text for text/PDF/DOCX/XLSX files (possibly empty).
        - `IMAGE_SENTINEL` for image files (caller routes to vision/OCR).
        - "" for unsupported types or on extraction failure.
    """
    ext = path.suffix.lower()

    if ext in IMAGE_EXTENSIONS or (mime or "").startswith("image/"):
        return IMAGE_SENTINEL

    try:
        if ext in TEXT_EXTENSIONS or (mime or "").startswith("text/"):
            return _truncate(_extract_text_file(path))
        if ext == ".pdf":
            return _truncate(_extract_pdf(path))
        if ext == ".docx":
            return _truncate(_extract_docx(path))
        if ext == ".xlsx":
            return _truncate(_extract_xlsx(path))
        # Legacy .xls is not supported by openpyxl — return "" so the handler
        # surfaces a friendly error rather than crashing.
        logger.info("Unsupported file type", path=str(path), ext=ext, mime=mime)
        return ""
    except Exception as e:
        logger.warning("File extraction failed", path=str(path), ext=ext, error=str(e))
        return ""
